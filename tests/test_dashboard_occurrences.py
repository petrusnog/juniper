"""Dashboard: ocorrências de vazamento e exportação em xlsx (T-168)."""

from __future__ import annotations

import io
import json
from collections.abc import AsyncIterator
from datetime import UTC, datetime, timedelta
from typing import Any

import mongomock
import openpyxl
import pytest

from brain.core.leak_audit import build_record, detect
from brain.core.leak_report import Exchange, build_report
from brain.dashboard.api import Api
from brain.dashboard.routes import build_api
from brain.dashboard.server import DashboardServer
from brain.dashboard.sources import Sources
from brain.dashboard.xlsx import MIME
from brain.memory.leak_audit import LeakAuditStore, LeakReportStore
from brain.memory.mongo import MongoConnection
from tests.test_dashboard_server import BEARER, TOKEN, get, login

SECRET = "sk-proj-abcdefghijklmnopqrstuvwxyz123456"
TEXT = "Segunda, com <anotacao_interna>vazada</anotacao_interna>."


def connection() -> MongoConnection:
    conn = MongoConnection()
    conn._client = mongomock.MongoClient(tz_aware=True)  # type: ignore[assignment]
    return conn


def make_report(**overrides: Any) -> dict[str, Any]:
    report = build_report(
        exchanges=[
            Exchange("primeira", "Primeira resposta.", "2026-10-29T10:00:00+00:00"),
            Exchange("segunda", TEXT, "2026-10-29T10:01:00+00:00"),
            Exchange("e a terceira", "Deixa eu ver", "2026-10-29T10:02:00+00:00", interrupted=True),
        ],
        session_id="20261029-100000-abc123",
        channel="telegram",
        voice=False,
        note="a tag apareceu na segunda",
        audit_ids=["f" * 32],
    )
    report.update(overrides)
    return report


def make_record(**overrides: Any) -> dict[str, Any]:
    raw = f"Pronto. <conhecimento>notas</conhecimento> chave {SECRET}"
    delivered = "Pronto. <anotacao>x</anotacao>"
    record = build_record(
        hits=detect(raw, delivered, [("persona", "Você é a Juniper.")], echo_words=10),
        guard_triggered=True,
        raw=raw,
        delivered=delivered,
        session_id="20261029-110000-def456",
        channel="mobile",
        voice=True,
        provider="deepinfra",
        model="m1",
        finish_reason="stop",
        tool_calls=[],
        blocks=[],
        user_message="[quinta-feira, 29/10/2026, 11:00 (America/Sao_Paulo)] me explica isso",
    )
    record.update(overrides)
    return record


class World:
    def __init__(self, *, reports: bool = True, audit: bool = True) -> None:
        conn = connection()
        self.reports = LeakReportStore(conn)
        self.audit = LeakAuditStore(conn, retention_days=30)
        self.api: Api = build_api(
            Sources(reports=self.reports if reports else None, audit=self.audit if audit else None)
        )

    async def get(self, path: str, **params: Any) -> tuple[int, Any]:
        status, body = await self.api.handle(path, {k: [str(v)] for k, v in params.items()})
        json.dumps(body)  # always JSON-able
        return status, body


@pytest.fixture
def world() -> World:
    return World()


async def test_the_list_merges_reports_and_detections_newest_first(world: World) -> None:
    old = make_report(at=datetime.now(UTC) - timedelta(hours=2))
    world.reports.add(old)
    world.audit.add(make_record())
    status, body = await world.get("/occurrences")
    assert status == 200
    assert [i["kind"] for i in body["items"]] == ["automatico", "relato"]
    detected, reported = body["items"]
    assert reported["note"] == "a tag apareceu na segunda"
    assert reported["session"] == "20261029-100000-abc123" and reported["channel"] == "telegram"
    assert reported["messages"] == 6  # three exchanges, the last one with her cut reply
    assert detected["voice"] is True and detected["rules"]
    assert len(reported["id"]) == 16  # short: a full 32-character id reads as a token


async def test_the_filters_pick_one_source(world: World) -> None:
    world.reports.add(make_report())
    world.audit.add(make_record())
    only_reports = (await world.get("/occurrences", kind="relato"))[1]
    only_detected = (await world.get("/occurrences", kind="automatico"))[1]
    assert [i["kind"] for i in only_reports["items"]] == ["relato"]
    assert [i["kind"] for i in only_detected["items"]] == ["automatico"]
    status, body = await world.get("/occurrences", kind="outro")
    assert status == 400 and "kind" in body["error"]


async def test_the_detail_of_a_report_has_the_messages_by_author(world: World) -> None:
    report = make_report()
    world.reports.add(report)
    short = report["_id"][:16]
    status, body = await world.get(f"/occurrences/{short}")
    assert status == 200 and body["kind"] == "relato"
    assert [(m["author"], m["text"]) for m in body["messages_list"]] == [
        ("Usuário", "primeira"),
        ("Juniper", "Primeira resposta."),
        ("Usuário", "segunda"),
        ("Juniper", TEXT),
        ("Usuário", "e a terceira"),
        ("Juniper", "Deixa eu ver"),
    ]
    assert body["messages_list"][-1]["note"] == "você cortou a resposta"
    assert body["audit"] == ["f" * 16]


async def test_the_detail_of_a_detection_has_what_was_written_and_what_was_shown(
    world: World,
) -> None:
    record = make_record()
    world.audit.add(record)
    status, body = await world.get(f"/occurrences/{record['_id'][:16]}")
    assert status == 200 and body["kind"] == "automatico"
    authors = [m["author"] for m in body["messages_list"]]
    assert authors == ["Usuário", "Juniper", "Juniper (texto bruto do modelo)"]
    assert body["messages_list"][0]["text"] == "me explica isso"  # without the date stamp
    assert body["hits"] and body["model"] == "m1" and body["leaked"] is True
    assert SECRET not in json.dumps(body)  # masked, like everything the dashboard sends


async def test_an_unknown_or_malformed_id_is_a_404(world: World) -> None:
    for bad in ("0123456789abcdef", "nao-e-hex", "zz"):
        status, _ = await world.get(f"/occurrences/{bad}")
        assert status == 404


async def test_without_mongodb_the_screen_says_what_is_missing() -> None:
    off = World(reports=False, audit=False)
    status, body = await off.get("/occurrences")
    assert status == 503 and "MongoDB" in body["error"]
    assert (await off.get("/occurrences/abcd"))[0] == 503
    status, body = await off.api.handle_file("/occurrences/export.xlsx", {})
    assert status == 503 and isinstance(body, dict)


async def test_with_only_one_store_the_other_kind_is_just_empty() -> None:
    only = World(audit=False)
    only.reports.add(make_report())
    status, body = await only.get("/occurrences", kind="todos")
    assert status == 200 and [i["kind"] for i in body["items"]] == ["relato"]


# --- the spreadsheet -----------------------------------------------------------------------------


def read_workbook(content: bytes) -> openpyxl.Workbook:
    return openpyxl.load_workbook(io.BytesIO(content))


async def test_the_export_has_one_row_per_message_with_session_author_and_text(
    world: World,
) -> None:
    report = make_report()
    world.reports.add(report)
    world.audit.add(make_record())
    status, body = await world.api.handle_file("/occurrences/export.xlsx", {})
    assert status == 200 and isinstance(body, tuple)
    name, kind, content = body
    assert name.startswith("ocorrencias-juniper-") and name.endswith(".xlsx") and kind == MIME

    book = read_workbook(content)
    assert book.sheetnames == ["Mensagens", "Ocorrências"]
    rows = list(book["Mensagens"].iter_rows(values_only=True))
    assert rows[0] == (
        "Ocorrência", "Tipo", "ID da sessão", "Canal", "Voz", "Data da ocorrência",
        "Ordem", "Autor", "Mensagem", "Observação",
    )  # fmt: skip
    mine = [r for r in rows[1:] if r[2] == "20261029-100000-abc123"]
    assert [(r[6], r[7], r[8]) for r in mine] == [
        (1, "Usuário", "primeira"),
        (2, "Juniper", "Primeira resposta."),
        (3, "Usuário", "segunda"),
        (4, "Juniper", TEXT),
        (5, "Usuário", "e a terceira"),
        (6, "Juniper", "Deixa eu ver"),
    ]
    assert all(r[1] == "Relato (/vazou)" and r[3] == "telegram" for r in mine)
    detected = [r for r in rows[1:] if r[2] == "20261029-110000-def456"]
    assert {r[7] for r in detected} == {"Usuário", "Juniper", "Juniper (texto bruto do modelo)"}
    assert SECRET not in "".join(str(c) for r in rows for c in r)

    summary = list(book["Ocorrências"].iter_rows(values_only=True))
    assert len(summary) == 3 and summary[0][2] == "ID da sessão"
    assert {r[2] for r in summary[1:]} == {"20261029-100000-abc123", "20261029-110000-def456"}


async def test_the_export_can_be_one_occurrence_or_one_kind(world: World) -> None:
    report = make_report()
    world.reports.add(report)
    world.audit.add(make_record())
    _, (_, _, only_one) = await world.api.handle_file(
        "/occurrences/export.xlsx", {"id": [report["_id"][:16]]}
    )  # type: ignore[misc]
    rows = list(read_workbook(only_one)["Ocorrências"].iter_rows(values_only=True))
    assert len(rows) == 2 and rows[1][1] == "Relato (/vazou)"
    _, (_, _, detected) = await world.api.handle_file(
        "/occurrences/export.xlsx", {"kind": ["automatico"]}
    )  # type: ignore[misc]
    rows = list(read_workbook(detected)["Ocorrências"].iter_rows(values_only=True))
    assert len(rows) == 2 and rows[1][1] == "Detectado"


async def test_a_message_that_looks_like_a_formula_stays_text(world: World) -> None:
    hostile = make_report(
        exchanges=[{"user": '=HYPERLINK("http://x","a")', "assistant": "@SUM(1)", "at": None}]
    )
    world.reports.add(hostile)
    _, (_, _, content) = await world.api.handle_file("/occurrences/export.xlsx", {})  # type: ignore[misc]
    cell = read_workbook(content)["Mensagens"]["I2"]
    assert cell.data_type == "s" and str(cell.value).startswith("=HYPERLINK")


# --- over HTTP -----------------------------------------------------------------------------------


@pytest.fixture
async def server() -> AsyncIterator[tuple[DashboardServer, World]]:
    world = World()
    dash = DashboardServer("127.0.0.1", 0, TOKEN, api=world.api)
    assert await dash.start()
    yield dash, world
    await dash.stop()


async def test_the_download_needs_the_login_and_comes_as_a_file(
    server: tuple[DashboardServer, World],
) -> None:
    dash, world = server
    world.reports.add(make_report())
    status, _, _ = await get(dash, "/api/occurrences/export.xlsx")
    assert status == 401

    cookie = await login(dash)
    status, headers, body = await get(
        dash, "/api/occurrences/export.xlsx?kind=relato", headers={"Cookie": cookie}
    )
    assert status == 200
    assert headers["content-type"] == MIME
    assert headers["content-disposition"].startswith('attachment; filename="ocorrencias-juniper-')
    assert headers["cache-control"] == "no-store" and headers["x-content-type-options"] == "nosniff"
    assert read_workbook(body).sheetnames == ["Mensagens", "Ocorrências"]

    status, _, body = await get(dash, "/api/occurrences/export.xlsx", headers=BEARER)
    assert status == 200 and body[:2] == b"PK"  # a zip


async def test_the_list_over_http_is_json(server: tuple[DashboardServer, World]) -> None:
    dash, world = server
    world.reports.add(make_report())
    status, headers, body = await get(dash, "/api/occurrences", headers=BEARER)
    assert status == 200 and headers["content-type"].startswith("application/json")
    assert json.loads(body)["items"][0]["kind"] == "relato"
